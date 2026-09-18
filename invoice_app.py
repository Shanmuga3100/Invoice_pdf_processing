import json
from io import BytesIO

import streamlit as st
from openpyxl import Workbook

# -------------------------------------------------------------------
# TODO: import / initialize your actual LLM object here.
# Example (LangChain + Anthropic):
# from langchain_anthropic import ChatAnthropic
# llm = ChatAnthropic(model="claude-sonnet-4-6", api_key="YOUR_API_KEY")
# Example (LangChain + OpenAI):
# from langchain_openai import ChatOpenAI
# llm = ChatOpenAI(model="gpt-4o", api_key="YOUR_API_KEY")
# -------------------------------------------------------------------
#llm =None   # <-- replace this with your real llm object
from dotenv import load_dotenv
from langchain_groq import ChatGroq
import os
import pdfplumber
load_dotenv()




llm = ChatGroq(
    model_name="openai/gpt-oss-120b",
    temperature=0
)


def extract_invoice_json(invoice_report: str) -> dict:
    """Send the raw invoice text to the LLM and parse the JSON response."""
    msg=f""" store the values in excel file and make it export.
INVOICE_REPORT:
make it as JSON format only the below values not any sentence
Format:
{{
  "Vendor Name":"",
  "Invoice Number":"",
  "Invoice Date":"",
  "Grand Total":""
  "Item":""
}} only this format value don;t add any additional sentence
{invoice_report}"""


    invoice_export = llm.invoke(msg)
    response = invoice_export.content
    response = response.replace("```json", "").replace("```", "").strip()
    data = json.loads(response)
    return data


def build_excel(data: dict) -> BytesIO:
    """Build an in-memory .xlsx file from the extracted key/value data."""
    

    wb=Workbook()
    ws=wb.active

    for row, (key, value) in enumerate(data.items(), start=1):
        ws.cell(row=row, column=1).value = key
        ws.cell(row=row, column=2).value = value
    wb.save("Invoice_pdf_report.xlsx")

    print("Saved successfully")

    # Auto-fit column width (simple heuristic)
    for col in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 5

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# -------------------------------------------------------------------
# Streamlit UI
# -------------------------------------------------------------------
st.set_page_config(page_title="Invoice to Excel Converter", page_icon="📄")
st.title("📄 Invoice PDF → Excel Converter")

st.write("Upload an invoice pdf file (.pdf). It will be processed and "
         "converted into a downloadable Excel report.")

uploaded_file = st.file_uploader("Select a pdf file", type=["pdf"])

if uploaded_file is not None:
    
    with pdfplumber.open(uploaded_file) as f:

        invoice_report=""
        for page in f.pages:
            invoice_report+=page.extract_text()
        f.close()
    #print(pdf_text)

    st.subheader("Preview of uploaded text")
    st.text_area("Invoice content", invoice_report, height=200)

    if st.button("Convert to Excel"):
        if llm is None:
            st.error("LLM is not configured. Please set up the `llm` object in app.py.")
        else:
            with st.spinner("Extracting invoice details..."):
                try:
                    data = extract_invoice_json(invoice_report)
                except json.JSONDecodeError:
                    st.error("Could not parse the model's response as JSON. "
                             "Please check the LLM output format.")
                    st.stop()
                except Exception as e:
                    st.error(f"Something went wrong: {e}")
                    st.stop()

            st.success("Invoice data extracted successfully!")
            st.json(data)

            excel_buffer = build_excel(data)

            st.download_button(
                label="⬇️ Download Excel",
                data=excel_buffer,
                file_name="Invoice_Report.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
else:
    st.info("Please upload a .pdf file to begin.")





